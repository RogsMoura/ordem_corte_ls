"""Gerador de Ordem de Corte para pedidos em PDF."""

from __future__ import annotations

import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pdfplumber
from PIL import Image, ImageTk
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAO_INFORMADO = "Nao informado"
PRODUCT_LINE = re.compile(
    r"^(?P<produto>.+?)\s+(?P<quantidade>\d+)\s+R\$\s*(?P<preco>[\d.,]+)\s*$",
    re.IGNORECASE,
)


def caminho_recurso(nome: str) -> Path:
    """Localiza recursos tanto no script quanto no executavel PyInstaller."""
    pasta = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return pasta / nome


@dataclass
class PedidoItem:
    pedido: str
    produto: str
    cor_estampa: str = NAO_INFORMADO
    cor_calca: str = NAO_INFORMADO
    genero: str = NAO_INFORMADO
    tipo_camisa: str = NAO_INFORMADO
    tecido_calca: str = NAO_INFORMADO
    tamanho_camisa: str = NAO_INFORMADO
    tamanho_calca: str = NAO_INFORMADO
    frisos_detalhes: list[str] = field(default_factory=list)
    tecido: str = NAO_INFORMADO
    modelagem: str = NAO_INFORMADO
    observacoes: list[str] = field(default_factory=list)


def normalizar(texto: str) -> str:
    """Remove acentos e deixa o texto comparavel sem alterar o original."""
    sem_acentos = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in sem_acentos if not unicodedata.combining(c)).casefold()


def normalizar_comparavel(texto: str) -> str:
    """Remove acentos, mas preserva quebras de linha e pontuacao."""
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def extrair_texto(pdf_path: Path) -> str:
    """Extrai texto de todas as paginas do PDF."""
    with pdfplumber.open(pdf_path) as pdf:
        return "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)


def encontrar_numero_pedido(texto: str) -> str:
    """Localiza o numero do pedido nos cabecalhos comuns do documento."""
    padrao = re.compile(r"numero\s+(?:do\s+)?pedido\s*:\s*([\w/-]+)", re.IGNORECASE)
    encontrado = padrao.search(normalizar_comparavel(texto))
    return encontrado.group(1) if encontrado else NAO_INFORMADO


def secoes_de_pedido(texto: str) -> list[str]:
    """Divide um PDF mesclado em uma seção por pedido do cliente."""
    linhas = texto.splitlines()
    secoes: list[list[str]] = []
    atual: list[str] = []
    encontrou_cabecalho = False
    for linha in linhas:
        if normalizar(linha) == "pedido":
            if atual and encontrou_cabecalho:
                secoes.append(atual)
            atual = [linha]
            encontrou_cabecalho = False
        elif atual:
            atual.append(linha)
            if re.search(r"numero\s+do\s+pedido\s*:", normalizar_comparavel(linha), re.IGNORECASE):
                encontrou_cabecalho = True
    if atual and encontrou_cabecalho:
        secoes.append(atual)
    return ["\n".join(secao) for secao in secoes]


def blocos_de_produto(linhas: Iterable[str]) -> list[list[str]]:
    """Separa o texto em blocos iniciados por uma linha de produto."""
    blocos: list[list[str]] = []
    atual: list[str] = []
    for linha in linhas:
        linha = linha.strip()
        if not linha:
            continue
        corresponde = PRODUCT_LINE.match(linha)
        if corresponde and not normalizar(linha).startswith("produto quantidade"):
            if atual:
                blocos.append(atual)
            atual = [linha]
        elif atual:
            atual.append(linha)
    if atual:
        blocos.append(atual)
    return blocos


def primeiro_valor(campos: dict[str, str], termos: tuple[str, ...]) -> str:
    for chave, valor in campos.items():
        if any(termo in normalizar(chave) for termo in termos) and valor:
            return valor
    return NAO_INFORMADO


def limpar_tamanho(valor: str) -> str:
    """Remove adicionais de preço que aparecem junto do tamanho no pedido."""
    return re.sub(r"\s*\(R\$[^)]*\)", "", valor, flags=re.IGNORECASE).strip()


def estampa_do_produto(produto: str) -> str:
    """Obtém uma descrição útil para a grade quando não há campo de cor."""
    descricao = re.sub(r"^pijama hospitalar\s+", "", produto, flags=re.IGNORECASE).strip()
    descricao = re.sub(r"^camisas estampadas\s*-\s*", "", descricao, flags=re.IGNORECASE).strip()
    descricao = re.sub(r"^com estampa\s+", "", descricao, flags=re.IGNORECASE).strip()
    return descricao or NAO_INFORMADO


def identificar_tecido(produto: str, campos: dict[str, str]) -> str:
    texto = normalizar(produto)
    if "estampad" in texto or "estampa" in texto:
        return "Estampado"
    tecidos = (("gabardine", "Gabardine"), ("oxfordine", "Oxfordine"), ("cedromix", "Cedromix"), ("melange", "Melange"), ("premium", "Premium"))
    for termo, nome in tecidos:
        if termo in texto:
            return nome
    return "Gabardine"


def identificar_tipo_camisa(produto: str, campos: dict[str, str]) -> str:
    """Define a modelagem da camisa usada na grade de corte."""
    texto = normalizar(" ".join([produto, *campos.keys()]))
    if any("tamanho camisa classic" in normalizar(chave) for chave in campos):
        return "Classic"
    if "babyl" in texto:
        return "BabyLook"
    return "Padrao"


def interpretar_bloco(bloco: list[str], numero_pedido: str) -> list[PedidoItem]:
    """Converte um bloco de produto em um registro estruturado."""
    produto_match = PRODUCT_LINE.match(bloco[0])
    if produto_match is None:
        raise ValueError(f"Linha de produto invalida: {bloco[0]}")
    produto = produto_match.group("produto").strip()
    campos: dict[str, str] = {}
    chave_atual = ""
    for linha in bloco[1:]:
        if ":" in linha:
            chave, valor = linha.split(":", 1)
            chave_atual = chave.strip()
            campos[chave_atual] = valor.strip()
        elif chave_atual and normalizar(chave_atual) == "observacoes":
            marcadores_fim = ("subtotal", "entrega", "total", "pac ", "sedex", "correios", "motoboy", "retirada")
            if normalizar(linha).startswith(marcadores_fim):
                chave_atual = ""
            else:
                campos[chave_atual] = f"{campos[chave_atual]} {linha.strip()}".strip()

    texto_campos = " ".join(f"{chave}: {valor}" for chave, valor in campos.items())
    genero = primeiro_valor(campos, ("escolha a tabela de tamanhos",))
    genero_normalizado = normalizar(genero)
    if "femin" in genero_normalizado:
        genero = "Feminino"
    else:
        genero = "Masculino"

    tamanho_camisa = limpar_tamanho(primeiro_valor(campos, ("tamanho parte de cima", "tamanho da camisa", "tamanho camisa")))
    tamanho_calca = limpar_tamanho(primeiro_valor(campos, ("tamanho da calca", "tamanho calca")))
    cor_conjunto = primeiro_valor(campos, ("cor do conjunto liso", "cor do conjunto"))
    estampa_kit = primeiro_valor(campos, ("estampas kits",))
    time_desejado = primeiro_valor(campos, ("time desejado",))
    cor = primeiro_valor(campos, ("cores", "estampa"))
    if cor_conjunto != NAO_INFORMADO:
        cor = cor_conjunto
    cor_calca = primeiro_valor(campos, ("cor da calca",))
    barra = primeiro_valor(campos, ("barra da calca", "barra"))
    deseja_frisos = primeiro_valor(campos, ("deseja adicionar frisos",))
    cor_frisos = primeiro_valor(campos, ("cor do friso", "cor dos frisos"))
    e_smile = primeiro_valor(campos, ("cores smile",)) != NAO_INFORMADO
    tipo_camisa = identificar_tipo_camisa(produto, campos)
    if tipo_camisa == "Classic":
        genero = f"{genero} Classic"
    detalhes = []
    if "sim" in normalizar(deseja_frisos) and cor_frisos != NAO_INFORMADO:
        detalhes.append(cor_frisos)
    detalhes.append("Smile" if e_smile else "X")
    observacao = primeiro_valor(campos, ("observacoes",))
    observacoes = [] if observacao == NAO_INFORMADO else [observacao]
    if cor == NAO_INFORMADO:
        cor = primeiro_valor(campos, ("camisa estampada", "estampas kits"))
    if cor == NAO_INFORMADO:
        cor = estampa_do_produto(produto)
    cor_base = cor
    if time_desejado != NAO_INFORMADO:
        cor = f"{time_desejado} - {cor}"
    if cor_calca == NAO_INFORMADO:
        cor_calca = cor_base
    if time_desejado != NAO_INFORMADO:
        cor_calca = f"{time_desejado} - {cor_calca}"
    modelagem = "Jogger" if "jogger" in normalizar(barra) else "Tradicional"

    item = PedidoItem(
        pedido=numero_pedido,
        produto=produto,
        cor_estampa=cor,
        cor_calca=cor_calca,
        genero=genero,
        tipo_camisa=tipo_camisa,
        tamanho_camisa=tamanho_camisa,
        tamanho_calca=tamanho_calca,
        frisos_detalhes=detalhes,
        tecido=identificar_tecido(produto, campos),
        tecido_calca="Gabardine" if identificar_tecido(produto, campos) == "Estampado" else identificar_tecido(produto, campos),
        modelagem=modelagem,
        observacoes=observacoes,
    )
    if "conjunto liso" in normalizar(produto) and estampa_kit != NAO_INFORMADO:
        item_estampado = PedidoItem(
            pedido=numero_pedido,
            produto=f"Camisa Estampada - {estampa_kit}",
            cor_estampa=estampa_kit,
            cor_calca=NAO_INFORMADO,
            genero=genero,
            tipo_camisa=tipo_camisa,
            tamanho_camisa=tamanho_camisa,
            tamanho_calca=NAO_INFORMADO,
            frisos_detalhes=detalhes.copy(),
            tecido="Estampado",
            tecido_calca="Gabardine",
            modelagem="Tradicional",
            observacoes=observacoes.copy(),
        )
        item.produto = "Conjunto Liso"
        item.cor_estampa = cor_conjunto
        item.cor_calca = cor_conjunto
        item.tecido = "Gabardine"
        return [item, item_estampado]
    return [item]


def interpretar_pedido(texto: str) -> list[PedidoItem]:
    itens: list[PedidoItem] = []
    secoes = secoes_de_pedido(texto)
    if not secoes:
        secoes = [texto]
    for secao in secoes:
        numero = encontrar_numero_pedido(secao)
        for bloco in blocos_de_produto(secao.splitlines()):
            itens.extend(interpretar_bloco(bloco, numero))
    return itens


def estilizar_tabela(planilha, linha_cabecalho: int, colunas: int) -> None:
    preenchimento = PatternFill("solid", fgColor="1F4E78")
    fonte_cabecalho = Font(color="FFFFFF", bold=True)
    borda = Border(*(Side(style="thin", color="B7C9D6") for _ in range(4)))
    for celula in planilha[linha_cabecalho]:
        if celula.column <= colunas:
            celula.fill = preenchimento
            celula.font = fonte_cabecalho
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celula.border = borda
    for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, max_col=colunas):
        for celula in linha:
            celula.border = borda
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def aplicar_fundo_por_pedido(planilha, primeira_linha: int, ultima_linha: int, colunas: int) -> None:
    """Alterna o fundo das linhas para destacar grupos de pedidos."""
    preenchimento_cinza = PatternFill("solid", fgColor="E7E6E6")
    pedido_atual = None
    grupo = -1
    for linha in range(primeira_linha, ultima_linha + 1):
        pedido = planilha.cell(linha, 1).value
        if pedido != pedido_atual:
            pedido_atual = pedido
            grupo += 1
        if grupo % 2 == 0:
            for coluna in range(1, colunas + 1):
                planilha.cell(linha, coluna).fill = preenchimento_cinza


def gerar_planilha(itens: list[PedidoItem], pdf_path: Path) -> Path:
    """Cria a planilha formatada ao lado do PDF original."""
    destino = pdf_path.with_suffix(".xlsx")
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Ordem de Corte"
    planilha.sheet_view.showGridLines = False

    titulo = planilha.cell(1, 1, "ORDEM DE CORTE - LS PIJAMAS")
    titulo.font = Font(size=16, bold=True, color="FFFFFF")
    titulo.fill = PatternFill("solid", fgColor="17365D")
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    planilha.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    planilha.row_dimensions[1].height = 28

    linha = 3
    colunas_camisa = ["NUMERO / PEDIDO", "COR / ESTAMPA", "GENERO", "TAMANHO", "FRISOS / DETALHES", "TECIDO", "OBSERVAÇÕES"]
    planilha.cell(linha, 1, "CAMISA").font = Font(size=13, bold=True, color="17365D")
    linha += 1
    for coluna, titulo_coluna in enumerate(colunas_camisa, 1):
        planilha.cell(linha, coluna, titulo_coluna)
    linha_cabecalho_camisa = linha
    linha += 1
    for item in itens:
        valores = [item.pedido, item.cor_estampa, item.genero, item.tamanho_camisa, " | ".join(item.frisos_detalhes) or NAO_INFORMADO, item.tecido, " ".join(item.observacoes) or "X"]
        for coluna, valor in enumerate(valores, 1):
            planilha.cell(linha, coluna, valor)
        linha += 1
    aplicar_fundo_por_pedido(planilha, linha_cabecalho_camisa + 1, linha - 1, len(colunas_camisa))
    estilizar_tabela(planilha, linha_cabecalho_camisa, len(colunas_camisa))

    linha += 2
    colunas_calca = ["NUMERO / PEDIDO", "COR", "TAMANHO", "BARRA", "TECIDO"]
    planilha.cell(linha, 1, "CALCA").font = Font(size=13, bold=True, color="17365D")
    linha += 1
    for coluna, titulo_coluna in enumerate(colunas_calca, 1):
        planilha.cell(linha, coluna, titulo_coluna)
    linha_cabecalho_calca = linha
    linha += 1
    for item in (item for item in itens if item.tamanho_calca != NAO_INFORMADO):
        valores = [item.pedido, item.cor_calca, item.tamanho_calca, item.modelagem, item.tecido_calca]
        for coluna, valor in enumerate(valores, 1):
            planilha.cell(linha, coluna, valor)
        linha += 1
    aplicar_fundo_por_pedido(planilha, linha_cabecalho_calca + 1, linha - 1, len(colunas_calca))
    estilizar_tabela(planilha, linha_cabecalho_calca, len(colunas_calca))

    for coluna in range(1, 7):
        letra = get_column_letter(coluna)
        maior = max((len(str(planilha.cell(l, coluna).value or "")) for l in range(1, planilha.max_row + 1)), default=10)
        planilha.column_dimensions[letra].width = min(max(maior + 3, 14), 42)
    planilha.column_dimensions["G"].width = 65
    for linha in range(linha_cabecalho_camisa + 1, linha_cabecalho_calca - 3):
        planilha.cell(linha, 7).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    planilha.freeze_panes = "A5"
    workbook.save(destino)
    return destino


class Aplicativo:
    def __init__(self, janela: tk.Tk):
        self.janela = janela
        self.pdf_path: Path | None = None
        janela.title("Gerador de Ordem de Corte - Ls Pijamas")
        self.logo = None
        caminho_icone = caminho_recurso("logo-ls.ico")
        if caminho_icone.exists():
            try:
                janela.iconbitmap(str(caminho_icone))
            except tk.TclError:
                pass
        janela.geometry("660x420")
        janela.minsize(560, 340)
        principal = ttk.Frame(janela, padding=20)
        principal.pack(fill="both", expand=True)
        caminho_logo = caminho_recurso("logo-ls.png")
        if caminho_logo.exists():
            try:
                imagem = Image.open(caminho_logo)
                imagem.thumbnail((180, 80))
                self.logo = ImageTk.PhotoImage(imagem)
                ttk.Label(principal, image=self.logo).pack(anchor="w", pady=(0, 8))
            except Exception:
                pass
        ttk.Label(principal, text="Gerador de Ordem de Corte", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        ttk.Label(principal, text="Selecione um pedido em PDF para gerar as tabelas de camisa e calca.").pack(anchor="w", pady=(4, 18))
        botoes = ttk.Frame(principal)
        botoes.pack(fill="x")
        self.botao_pdf = ttk.Button(botoes, text="Selecionar PDF do Pedido", command=self.selecionar_pdf)
        self.botao_pdf.pack(side="left")
        self.botao_gerar = ttk.Button(botoes, text="Gerar Planilha de Corte", command=self.gerar, state="disabled")
        self.botao_gerar.pack(side="left", padx=(10, 0))
        self.arquivo_label = ttk.Label(principal, text="Nenhum PDF selecionado", foreground="#555555")
        self.arquivo_label.pack(anchor="w", pady=12)
        self.log = tk.Text(principal, height=12, state="disabled", wrap="word", relief="solid", borderwidth=1)
        self.log.pack(fill="both", expand=True)

    def registrar(self, mensagem: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", mensagem + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def selecionar_pdf(self) -> None:
        escolhido = filedialog.askopenfilename(title="Selecione o PDF do pedido", filetypes=[("Arquivos PDF", "*.pdf"), ("Todos os arquivos", "*.*")])
        if escolhido:
            self.pdf_path = Path(escolhido)
            self.arquivo_label.configure(text=str(self.pdf_path))
            self.botao_gerar.configure(state="normal")
            self.registrar(f"PDF selecionado: {self.pdf_path.name}")

    def gerar(self) -> None:
        if self.pdf_path is None:
            return
        try:
            self.registrar("Lendo o PDF e identificando os itens...")
            itens = interpretar_pedido(extrair_texto(self.pdf_path))
            if not itens:
                raise ValueError("Nenhum item de produto foi identificado no PDF.")
            destino = gerar_planilha(itens, self.pdf_path)
            mensagem = f"Planilha gerada com sucesso: {destino} ({len(itens)} itens de camisa)."
            self.registrar(mensagem)
            messagebox.showinfo("Concluido", mensagem)
        except Exception as erro:
            self.registrar(f"Erro: {erro}")
            messagebox.showerror("Nao foi possivel gerar a planilha", str(erro))


def main() -> None:
    janela = tk.Tk()
    try:
        ttk.Style(janela).theme_use("clam")
    except tk.TclError:
        pass
    Aplicativo(janela)
    janela.mainloop()


if __name__ == "__main__":
    main()